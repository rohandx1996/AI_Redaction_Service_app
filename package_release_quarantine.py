"""Package source and redacted PDFs according to the release audit workbook."""

from __future__ import annotations

import argparse
import json
import shutil
from pathlib import Path

from openpyxl import load_workbook


def audit_decisions(audit_path: Path) -> dict[str, str]:
    workbook = load_workbook(audit_path, read_only=True, data_only=True)
    sheet = workbook["Summary"]
    rows = sheet.iter_rows(min_row=3, values_only=True)
    decisions = {}
    for pdf_path, decision, *_ in rows:
        if not pdf_path:
            continue
        decisions[Path(str(pdf_path)).name] = str(decision or "FAIL").upper()
    workbook.close()
    return decisions


def package_batch(batch_dir: Path) -> tuple[int, int, int]:
    audit_path = batch_dir / "release_audit.xlsx"
    if not audit_path.exists():
        raise FileNotFoundError(f"Audit workbook not found: {audit_path}")

    decisions = audit_decisions(audit_path)
    manifest = json.loads((batch_dir / "manifest.json").read_text(encoding="utf-8"))
    records = {item["document"]: item for item in manifest["documents"]}
    packages = []
    for document_dir in sorted(batch_dir.iterdir()):
        info_path = document_dir / "redaction_information.json"
        if not document_dir.is_dir() or not info_path.exists():
            continue
        record = records.get(document_dir.name)
        if not record:
            continue
        info = json.loads(info_path.read_text(encoding="utf-8"))
        source_pdf = Path(str(info.get("source_pdf", "")))
        redacted_path = Path(record.get("release_file") or record.get("quarantine_file") or "")
        packages.append((document_dir.name, source_pdf, redacted_path))

    existing_outputs = {}
    for record in packages:
        redacted_path = record[2]
        if redacted_path.exists():
            existing_outputs[redacted_path.name] = redacted_path

    for folder_name in ("release", "quarantine"):
        folder = batch_dir / folder_name
        folder.mkdir(parents=True, exist_ok=True)

    packaged = {"release": 0, "quarantine": 0}
    missing = 0
    for document_key, source_pdf, redacted_path in packages:
        decision = decisions.get(redacted_path.name, "FAIL")
        destination_name = "release" if decision == "PASS" else "quarantine"
        destination = batch_dir / destination_name / document_key
        destination.mkdir(parents=True)
        if source_pdf.exists():
            shutil.copy2(source_pdf, destination / source_pdf.name)
        else:
            missing += 1
        redacted_path = redacted_path if redacted_path.exists() else existing_outputs.get(redacted_path.name, Path())
        if redacted_path.is_file():
            shutil.copy2(redacted_path, destination / redacted_path.name)
        else:
            missing += 1
        packaged[destination_name] += 1

    return packaged["release"], packaged["quarantine"], missing


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("batch_dir", type=Path)
    args = parser.parse_args()
    release, quarantine, missing = package_batch(args.batch_dir)
    print(f"release={release} quarantine={quarantine} missing_files={missing}")


if __name__ == "__main__":
    main()
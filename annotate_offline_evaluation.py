"""Annotate a consolidated redaction batch with deterministic offline accuracy results."""

from __future__ import annotations

import argparse
import json
import re
import shutil
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from compare_redaction_ground_truth import load_ground_truth, match_entities, normalize, pred_group, source_key

STRONG_PROTECTED_CATEGORIES = {
    "SUPPLIER_NAME", "PARENT_COMPANY_NAME", "MANUFACTURER_NAME", "SUBCONTRACTOR_NAME",
    "COMMERCIAL_PARTNER_NAME", "CUSTOMER_NAME", "PERSON_NAME", "EMAIL", "PHONE", "FAX",
    "ADDRESS", "PLANT_NAME", "SITE_NAME", "SIGNATURE_TEXT", "SUPPLIER_CODE",
    "SUPPLIER_IDENTIFIER", "CUSTOMER_IDENTIFIER", "REPORT_NUMBER", "CERTIFICATE_NUMBER",
    "DOCUMENT_ID", "PIN_CODE", "INTERNAL_REFERENCE", "NDA_REFERENCE", "CONTRACT_REFERENCE",
    "QR_CODE_TEXT", "BARCODE_TEXT",
}
CONTEXTUAL_PROTECTED_CATEGORIES = {
    "COMPANY", "CONTACT_ROLE", "LOCATION", "SAMPLE_NAME", "PROJECT_NAME", "PRODUCT_NAME",
    "PRODUCT_CODE", "PART_NUMBER", "MATERIAL_ID", "BOM_ID", "FORMULATION_NAME", "TRADE_NAME",
    "PROPRIETARY_MATERIAL_DATA", "URL",
}
PROTECTED_LABELS = {
    "client name", "customer", "applicant", "submitted by", "sample submitted by", "client address",
    "sample name", "sample description", "sample identification", "style/item", "model no",
    "part number", "client ref", "customer ref", "order no", "lot no", "sample no", "sample id",
    "report no", "certificate no", "pin code", "manufacturer", "supplier", "contact", "address",
}
ROUTINE_LAB_TERMS = {"sgs", "intertek", "bureau veritas", "tuv", "testing laboratory", "laboratory"}
GENERIC_LABEL_VALUES = {normalize(value) for value in PROTECTED_LABELS}


def prediction_records(document: dict) -> tuple[list[dict], list[tuple[dict, dict]]]:
    records = []
    references = []
    for page in document.get("pages", []) or []:
        page_number = int(page.get("page_number", 0)) + 1
        for category, entities in (page.get("entities", {}) or {}).items():
            for entity in entities or []:
                value = str(entity.get("value", "")).strip()
                if not value:
                    continue
                record = {
                    "document": source_key(document.get("source_pdf", "")),
                    "page": page_number,
                    "category": category,
                    "group": pred_group(category, value),
                    "value": value,
                    "norm_value": normalize(value),
                    "span_indices": entity.get("span_indices", []),
                    "confidence": entity.get("confidence"),
                }
                records.append(record)
                references.append((entity, record))
    return records, references


def ratio(numerator: int, denominator: int) -> float:
    return numerator / denominator if denominator else 0.0


def retained_values(path: Path) -> dict[str, str]:
    data = json.loads(path.read_text(encoding="utf-8"))
    values = {}
    for category, entries in data.get("unique_retained_knowledge", {}).items():
        for value in entries or []:
            normalized = normalize(value)
            if normalized:
                values[normalized] = category
    return values


def page_context(recognition: dict) -> dict[int, dict[int, str]]:
    result = {}
    for page in recognition.get("pages", []) or []:
        spans = page.get("spans", []) or []
        result[int(page.get("page_number", 0)) + 1] = {
            int(span.get("index", index)): str(span.get("text", ""))
            for index, span in enumerate(spans)
        }
    return result


def nearby_text(context: dict[int, dict[int, str]], page: int, indices: list[int]) -> str:
    spans = context.get(page, {})
    nearby = set()
    for index in indices:
        nearby.update(range(int(index) - 2, int(index) + 3))
    return " ".join(spans[index] for index in sorted(nearby) if index in spans).casefold()


def policy_classification(record: dict, context: str, retained: dict[str, str]) -> tuple[str, str]:
    category = record["category"].upper()
    normalized = record["norm_value"]
    value = record["value"]
    matched_labels = sorted(label for label in PROTECTED_LABELS if label in context)
    if normalized in GENERIC_LABEL_VALUES:
        return "LIKELY_FALSE_POSITIVE", "The extracted value is a generic field label, not its value"
    if normalized in retained:
        return "LIKELY_FALSE_POSITIVE", f"Exact match in published retained-knowledge inventory: {retained[normalized]}"
    if category == "LABORATORY_NAME" and any(term in value.casefold() or term in context for term in ROUTINE_LAB_TERMS):
        return "LIKELY_FALSE_POSITIVE", "Routine testing-laboratory identity is retained by policy"
    if re.fullmatch(r"\d{2,7}-\d{2}-\d", value.strip()) or re.search(r"\b(?:iec|iso|astm|en|epa)\s*\d", value, re.I):
        return "LIKELY_FALSE_POSITIVE", "Value resembles retained CAS/standard technical knowledge"
    if matched_labels:
        return "POLICY_SUPPORTED", f"Nearby protected field context: {', '.join(matched_labels)}"
    if category in STRONG_PROTECTED_CATEGORIES:
        return "POLICY_SUPPORTED", f"Category {category} is explicitly protected by the internal policy"
    if category in CONTEXTUAL_PROTECTED_CATEGORIES:
        return "REVIEW_REQUIRED", f"Category {category} requires stronger local role or proprietary context"
    return "REVIEW_REQUIRED", "No deterministic policy rule establishes whether this value should be redacted"


def annotate_document(
    document: dict,
    truth: list[dict] | None,
    recognition: dict,
    retained: dict[str, str],
) -> dict:
    predictions, references = prediction_records(document)
    if truth is None:
        context_by_page = page_context(recognition)
        counts = defaultdict(int)
        for entity, record in references:
            context = nearby_text(context_by_page, record["page"], record.get("span_indices", []))
            classification, evidence = policy_classification(record, context, retained)
            entity["offline_evaluation"] = classification
            entity["offline_evaluation_basis"] = "OFFLINE_POLICY_REVIEW"
            entity["offline_evaluation_evidence"] = evidence
            counts[classification] += 1
        return {
            "status": "POLICY_EVALUATED",
            "reason": "No matched ground truth; predictions reviewed with internal policy and published inventories",
            "predicted_entities": len(predictions),
            "policy_supported": counts["POLICY_SUPPORTED"],
            "likely_false_positives": counts["LIKELY_FALSE_POSITIVE"],
            "review_required": counts["REVIEW_REQUIRED"],
            "policy_support_rate": ratio(counts["POLICY_SUPPORTED"], len(predictions)),
            "accuracy_limitation": "False negatives, recall, and true precision cannot be measured without ground truth",
        }

    matches, missed, false_positives = match_entities(truth, predictions)
    matched_ids = {id(item["prediction"]): item for item in matches}
    false_positive_ids = {id(item) for item in false_positives}
    for entity, record in references:
        match = matched_ids.get(id(record))
        if match:
            entity["offline_evaluation"] = "TRUE_POSITIVE"
            entity["ground_truth_category"] = match["truth"]["category"]
            entity["ground_truth_field_label"] = match["truth"].get("field_label", "")
        elif id(record) in false_positive_ids:
            entity["offline_evaluation"] = "FALSE_POSITIVE"

    true_positives = len(matches)
    return {
        "status": "GROUND_TRUTH_EVALUATED",
        "matching_rule": "One-to-one exact normalized value on the same 1-based page",
        "ground_truth_entities": len(truth),
        "predicted_entities": len(predictions),
        "true_positives": true_positives,
        "false_negatives": len(missed),
        "false_positives": len(false_positives),
        "recall": ratio(true_positives, len(truth)),
        "precision": ratio(true_positives, len(predictions)),
        "f1": ratio(2 * true_positives, len(truth) + len(predictions)),
        "false_negative_details": [
            {
                "page": item["page"],
                "category": item["category"],
                "value": item["value"],
                "field_label": item.get("field_label", ""),
            }
            for item in missed
        ],
    }


def add_audit_sheets(audit_path: Path, output_path: Path, documents: list[dict]) -> None:
    shutil.copy2(audit_path, output_path)
    workbook = load_workbook(output_path)
    for name in ("Offline Accuracy", "Offline Errors"):
        if name in workbook.sheetnames:
            del workbook[name]

    accuracy = workbook.create_sheet("Offline Accuracy")
    headers = [
        "Document", "Source PDF", "Evaluation Status", "Ground Truth", "Predicted",
        "True Positives", "False Negatives", "False Positives", "Policy Supported",
        "Likely False Positives", "Review Required", "Recall", "Precision", "F1", "Policy Support Rate",
    ]
    accuracy.append(headers)
    errors = workbook.create_sheet("Offline Errors")
    errors.append(["Document", "Page", "Error Type", "Category", "Value", "Field Label"])
    for document in documents:
        evaluation = document["offline_evaluation"]
        accuracy.append([
            document.get("document", ""), document.get("source_pdf", ""), evaluation["status"],
            evaluation.get("ground_truth_entities", ""), evaluation.get("predicted_entities", 0),
            evaluation.get("true_positives", ""), evaluation.get("false_negatives", ""),
            evaluation.get("false_positives", ""), evaluation.get("policy_supported", ""),
            evaluation.get("likely_false_positives", ""), evaluation.get("review_required", ""),
            evaluation.get("recall", ""), evaluation.get("precision", ""), evaluation.get("f1", ""),
            evaluation.get("policy_support_rate", ""),
        ])
        for item in evaluation.get("false_negative_details", []):
            errors.append([
                document.get("document", ""), item["page"], "FALSE_NEGATIVE", item["category"],
                item["value"], item.get("field_label", ""),
            ])
        for page in document.get("pages", []) or []:
            for category, entities in (page.get("entities", {}) or {}).items():
                for entity in entities or []:
                    if entity.get("offline_evaluation") == "FALSE_POSITIVE":
                        errors.append([
                            document.get("document", ""), int(page.get("page_number", 0)) + 1,
                            "FALSE_POSITIVE", category, entity.get("value", ""), "",
                        ])
                    elif entity.get("offline_evaluation") in {"LIKELY_FALSE_POSITIVE", "REVIEW_REQUIRED"}:
                        errors.append([
                            document.get("document", ""), int(page.get("page_number", 0)) + 1,
                            entity["offline_evaluation"], category, entity.get("value", ""),
                            entity.get("offline_evaluation_evidence", ""),
                        ])

    header_fill = PatternFill("solid", fgColor="1F6FB2")
    for sheet in (accuracy, errors):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        for column in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 55)
            sheet.column_dimensions[column[0].column_letter].width = max(width, 12)
    workbook.save(output_path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--batch", type=Path, required=True)
    parser.add_argument("--ground-truth", type=Path, required=True)
    parser.add_argument("--retained-knowledge", type=Path, required=True)
    args = parser.parse_args()

    consolidated_path = args.batch / "consolidated_redaction.json"
    consolidated = json.loads(consolidated_path.read_text(encoding="utf-8"))
    ground_truth = load_ground_truth(args.ground_truth)
    retained = retained_values(args.retained_knowledge)
    counts = defaultdict(int)
    for document in consolidated.get("documents", []):
        truth = ground_truth.get(source_key(document.get("source_pdf", "")))
        recognition_path = args.batch / document.get("document", "") / "recognition.json"
        recognition = json.loads(recognition_path.read_text(encoding="utf-8")) if recognition_path.exists() else {}
        document["offline_evaluation"] = annotate_document(document, truth, recognition, retained)
        counts[document["offline_evaluation"]["status"]] += 1

    evaluated_json = args.batch / "consolidated_redaction_evaluated.json"
    evaluated_json.write_text(json.dumps(consolidated, ensure_ascii=False, indent=2), encoding="utf-8")
    evaluated_audit = args.batch / "release_audit_with_offline_accuracy.xlsx"
    add_audit_sheets(args.batch / "release_audit.xlsx", evaluated_audit, consolidated["documents"])
    print(json.dumps({
        "ground_truth_evaluated_documents": counts["GROUND_TRUTH_EVALUATED"],
        "policy_evaluated_documents": counts["POLICY_EVALUATED"],
        "json": str(evaluated_json),
        "audit": str(evaluated_audit),
    }, indent=2))


if __name__ == "__main__":
    main()
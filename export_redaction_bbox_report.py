"""Export release/quarantine redaction bboxes sorted by normalized area."""

from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image


def bbox_row(
    document: str,
    disposition: str,
    source_pdf: str,
    page: dict,
    source_type: str,
    category: str,
    entity: dict,
    image_width: int,
    image_height: int,
) -> dict | None:
    bbox = entity.get("bbox") or {}
    try:
        x0 = float(bbox["x0"])
        y0 = float(bbox["y0"])
        x1 = float(bbox["x1"])
        y1 = float(bbox["y1"])
    except (KeyError, TypeError, ValueError):
        return None
    width = max(0.0, x1 - x0)
    height = max(0.0, y1 - y0)
    area = width * height
    page_area = image_width * image_height
    return {
        "Document": document,
        "Disposition": disposition,
        "Source PDF": source_pdf,
        "Page": int(page.get("page_number", 0)) + 1,
        "Page Label": page.get("page_label", ""),
        "Source Type": source_type,
        "Information Label": category,
        "Information": entity.get("value", entity.get("text", "")),
        "Confidence": entity.get("confidence", ""),
        "Span Indices": ", ".join(str(value) for value in entity.get("span_indices", [])),
        "X0": x0,
        "Y0": y0,
        "X1": x1,
        "Y1": y1,
        "Width (px)": width,
        "Height (px)": height,
        "Area (px^2)": area,
        "Page Width (px)": image_width,
        "Page Height (px)": image_height,
        "Page Area (%)": area / page_area if page_area else 0.0,
    }


def collect(batch_dir: Path) -> tuple[list[dict], list[dict]]:
    manifest = json.loads((batch_dir / "manifest.json").read_text(encoding="utf-8"))
    statuses = {item["document"]: item for item in manifest.get("documents", [])}
    rows = []
    unlocated = []
    for document, record in statuses.items():
        if record.get("status") not in {"OK", "RELEASE_FAILED"}:
            continue
        disposition = "RELEASE" if record["status"] == "OK" else "QUARANTINE"
        package_dir = batch_dir / document
        recognition_path = package_dir / "recognition.json"
        if not recognition_path.exists():
            continue
        recognition = json.loads(recognition_path.read_text(encoding="utf-8"))
        for page in recognition.get("pages", []) or []:
            page_label = page.get("page_label", "")
            image_path = package_dir.parent.parent / "batch" / "paddle_dbnet_svtr_pages" / document / page_label / f"{page_label}.png"
            if not image_path.exists():
                image_path = Path(__file__).resolve().parent / "output" / "batch" / "paddle_dbnet_svtr_pages" / document / page_label / f"{page_label}.png"
            if not image_path.exists():
                continue
            with Image.open(image_path) as image:
                image_width, image_height = image.size
            for category, entities in (page.get("entities", {}) or {}).items():
                for entity in entities or []:
                    row = bbox_row(
                        document, disposition, recognition.get("source_pdf", ""), page,
                        "OCR_ENTITY", category, entity, image_width, image_height,
                    )
                    if row:
                        rows.append(row)
            for entity in page.get("visual_redactions", []) or []:
                row = bbox_row(
                    document, disposition, recognition.get("source_pdf", ""), page,
                    "VISUAL_ENTITY", str(entity.get("category", "OTHER_IDENTITY")), entity,
                    image_width, image_height,
                )
                if row:
                    rows.append(row)
        for detection in (record.get("release_audit") or {}).get("machine_readable_detections", []) or []:
            unlocated.append({
                "Document": document,
                "Disposition": disposition,
                "Page": int(detection.get("page_number", 0)) + 1,
                "Type": detection.get("type", "Unknown"),
                "Decoded": bool(detection.get("decoded")),
                "Reason": "Audit record does not preserve the detected bbox",
            })
    rows.sort(key=lambda row: (row["Page Area (%)"], row["Area (px^2)"]), reverse=True)
    for rank, row in enumerate(rows, 1):
        row["Area Rank"] = rank
    return rows, unlocated


def write_sheet(workbook: Workbook, name: str, rows: list[dict], headers: list[str]) -> None:
    sheet = workbook.create_sheet(name)
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])
    fill = PatternFill("solid", fgColor="1F6FB2")
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_index, cells in enumerate(sheet.columns, 1):
        width = min(max(len(str(cell.value or "")) for cell in cells) + 2, 55)
        sheet.column_dimensions[get_column_letter(column_index)].width = max(width, 12)
    if "Page Area (%)" in headers:
        column = get_column_letter(headers.index("Page Area (%)") + 1)
        for cell in sheet[column][1:]:
            cell.number_format = "0.0000%"


def export(batch_dir: Path, output_path: Path) -> None:
    rows, unlocated = collect(batch_dir)
    workbook = Workbook()
    workbook.remove(workbook.active)
    headers = [
        "Area Rank", "Document", "Disposition", "Page", "Page Label", "Source Type",
        "Information Label", "Information", "Confidence", "Span Indices", "X0", "Y0", "X1", "Y1",
        "Width (px)", "Height (px)", "Area (px^2)", "Page Width (px)", "Page Height (px)",
        "Page Area (%)", "Source PDF",
    ]
    write_sheet(workbook, "All Bboxes", rows, headers)
    write_sheet(workbook, "Visual Bboxes", [row for row in rows if row["Source Type"] == "VISUAL_ENTITY"], headers)
    write_sheet(workbook, "OCR Bboxes", [row for row in rows if row["Source Type"] == "OCR_ENTITY"], headers)
    write_sheet(workbook, "Release Bboxes", [row for row in rows if row["Disposition"] == "RELEASE"], headers)
    write_sheet(workbook, "Quarantine Bboxes", [row for row in rows if row["Disposition"] == "QUARANTINE"], headers)
    summary_counts = Counter((row["Disposition"], row["Information Label"]) for row in rows)
    summary = [
        {"Disposition": disposition, "Information Label": label, "BBox Count": count}
        for (disposition, label), count in sorted(summary_counts.items())
    ]
    write_sheet(workbook, "Summary", summary, ["Disposition", "Information Label", "BBox Count"])
    write_sheet(
        workbook, "QR Barcode No BBox", unlocated,
        ["Document", "Disposition", "Page", "Type", "Decoded", "Reason"],
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    print(json.dumps({
        "bbox_rows": len(rows),
        "release_rows": sum(row["Disposition"] == "RELEASE" for row in rows),
        "quarantine_rows": sum(row["Disposition"] == "QUARANTINE" for row in rows),
        "unlocated_machine_readable": len(unlocated),
        "output": str(output_path),
    }, indent=2))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("batch_dir", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    output = args.output or args.batch_dir / "redaction_bbox_area_report.xlsx"
    export(args.batch_dir, output)


if __name__ == "__main__":
    main()
"""Write internal PDF release-audit workbooks without exposing sensitive values."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


TITLE_FILL = PatternFill("solid", fgColor="0B2C4A")
HEADER_FILL = PatternFill("solid", fgColor="1F6FB2")
PASS_FILL = PatternFill("solid", fgColor="C6EFCE")
FAIL_FILL = PatternFill("solid", fgColor="FFC7CE")
WHITE_BOLD = Font(color="FFFFFF", bold=True)
THIN_BORDER = Border(*(Side(style="thin", color="C9D6E3") for _ in range(4)))


def _write_sheet(workbook: Workbook, name: str, title: str, headers: list[str], rows: list[list[object]]) -> None:
    sheet = workbook.create_sheet(name)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = sheet.cell(1, 1, title)
    title_cell.fill = TITLE_FILL
    title_cell.font = Font(color="FFFFFF", bold=True, size=14)
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(2, column, header)
        cell.fill = HEADER_FILL
        cell.font = WHITE_BOLD
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER
    for row_index, row in enumerate(rows, 3):
        for column, value in enumerate(row, 1):
            cell = sheet.cell(row_index, column, value)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = THIN_BORDER
            if value == "PASS":
                cell.fill = PASS_FILL
            elif value == "FAIL":
                cell.fill = FAIL_FILL
    for column in range(1, len(headers) + 1):
        width = max(len(str(headers[column - 1])), *(len(str(row[column - 1])) for row in rows)) if rows else len(str(headers[column - 1]))
        sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 12), 45)
    sheet.freeze_panes = "A3"


def write_release_audit(output_path: str | Path, documents: list[dict]) -> Path:
    """Create a release-audit workbook from batch document audit records."""
    output_file = Path(output_path)
    output_file.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    workbook.remove(workbook.active)

    summary_rows = []
    object_rows = []
    machine_readable_rows = []
    timing_rows = []
    for document in documents:
        audit = document.get("release_audit", {})
        findings = audit.get("findings", {})
        summary_rows.append([
            document.get("release_file", ""), audit.get("status", "FAIL"),
            audit.get("page_count", 0), audit.get("sensitive_value_count", 0),
            len(findings.get("leaked_values", [])),
        ])
        for check, value in findings.items():
            status = "PASS" if not value else "FAIL"
            if isinstance(value, (dict, list)):
                value = len(value)
            object_rows.append([document.get("release_file", ""), check, value, status])
        for detection in audit.get("machine_readable_detections", []):
            machine_readable_rows.append([
                document.get("release_file", ""), detection.get("page_number", 0) + 1,
                detection.get("type", "Unknown"), "Yes" if detection.get("decoded") else "No", "REDACTED",
            ])
        timings = document.get("timings_sec", {})
        timing_rows.append([
            document.get("release_file", "") or document.get("quarantine_file", ""),
            timings.get("qr_barcode_detection", 0), timings.get("pdf_redaction_and_audit", 0),
            timings.get("document_total", 0),
        ])

    _write_sheet(
        workbook, "Summary", "PDF Release Audit Summary",
        ["Release PDF", "Decision", "Pages", "Selected Values", "Text Leaks"], summary_rows,
    )
    _write_sheet(
        workbook, "PDF Objects", "Final PDF Object Checks",
        ["Release PDF", "Check", "Remaining Count", "Status"], object_rows,
    )
    _write_sheet(
        workbook, "QR and Barcodes", "Machine-Readable Marks Automatically Redacted",
        ["Release PDF", "Page", "Type", "Decoded", "Action"], machine_readable_rows,
    )
    _write_sheet(
        workbook, "Timings", "Document Processing Times (Seconds)",
        ["Output PDF", "QR/Barcode", "PDF Redaction + Audit", "Document Total"], timing_rows,
    )
    workbook.save(output_file)
    return output_file
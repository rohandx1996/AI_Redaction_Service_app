"""Build a ground-truth Excel workbook (entity category + redaction value per document)
from the aggregated original-vs-published redaction extraction JSON.

Source: output/matched_original_counterparts/<batch>/litellm_all_original_vs_published_redactions.json
Output: output/matched_original_counterparts/<batch>/ground_truth_redactions.xlsx
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

TITLE_FILL = PatternFill("solid", fgColor="0B2C4A")
HEADER_FILL = PatternFill("solid", fgColor="1F6FB2")
BAND_FILL = PatternFill("solid", fgColor="EAF1FA")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, color="FFFFFF", size=14)
THIN = Side(style="thin", color="B9C7D6")
THIN_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def document_label(original_pdf: str) -> str:
    stem = Path(original_pdf).stem
    return re.sub(r"_Original$", "", stem, flags=re.IGNORECASE)


def load_records(agg_path: Path) -> tuple[list[dict], list[dict]]:
    data = json.loads(agg_path.read_text(encoding="utf-8"))
    records: list[dict] = []
    doc_summaries: list[dict] = []
    for doc in data.get("documents", []):
        label = document_label(doc.get("original_pdf", ""))
        items = doc.get("redacted_information", []) or []
        for item in items:
            records.append(
                {
                    "Document": label,
                    "Original PDF": doc.get("original_pdf", ""),
                    "Published PDF": doc.get("published_pdf", ""),
                    "Page": item.get("page_number"),
                    "Entity Category": item.get("category", ""),
                    "Field Label": item.get("field_label", ""),
                    "Ground Truth Value (Original)": item.get("original_value", ""),
                    "Published State": item.get("published_state", ""),
                    "Published Visible Value": item.get("published_visible_value", ""),
                    "Confidence": item.get("confidence"),
                    "Evidence": item.get("evidence", ""),
                }
            )
        doc_summaries.append(
            {
                "Document": label,
                "Page Count": doc.get("page_count"),
                "Redacted Item Count": doc.get("redacted_item_count", len(items)),
                "Error Count": len(doc.get("errors") or []),
            }
        )
    return records, doc_summaries


def style_sheet(ws, df: pd.DataFrame, title: str) -> None:
    ncols = len(df.columns)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(ncols, 1))
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.fill = TITLE_FILL
    title_cell.font = TITLE_FONT
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 24

    header_row = 2
    for col_idx, col_name in enumerate(df.columns, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=col_name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER

    for row_idx, row in enumerate(df.itertuples(index=False), start=header_row + 1):
        for col_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if (row_idx - header_row) % 2 == 0:
                cell.fill = BAND_FILL

    for col_idx, col_name in enumerate(df.columns, start=1):
        max_len = max([len(str(col_name))] + [len(str(v)) for v in df[col_name].astype(str).tolist()])
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(12, max_len + 2), 60)

    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    last_col_letter = get_column_letter(ncols)
    last_row = header_row + len(df)
    if last_row > header_row:
        table = Table(
            displayName=re.sub(r"\W+", "_", title)[:30] or "Table1",
            ref=f"A{header_row}:{last_col_letter}{last_row}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showRowStripes=False, showFirstColumn=False
        )
        ws.add_table(table)


def build_workbook(agg_path: Path, output_path: Path) -> None:
    records, doc_summaries = load_records(agg_path)
    detail_df = pd.DataFrame(records)
    summary_df = pd.DataFrame(doc_summaries)

    category_df = (
        detail_df.groupby("Entity Category", dropna=False)
        .size()
        .reset_index(name="Occurrence Count")
        .sort_values("Occurrence Count", ascending=False)
        .reset_index(drop=True)
    )

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Document Summary", index=False, startrow=1)
        category_df.to_excel(writer, sheet_name="Category Summary", index=False, startrow=1)
        detail_df.to_excel(writer, sheet_name="Ground Truth Detail", index=False, startrow=1)

        wb = writer.book
        style_sheet(wb["Document Summary"], summary_df, "Ground Truth - Per Document Summary")
        style_sheet(wb["Category Summary"], category_df, "Ground Truth - Entity Category Distribution")
        style_sheet(wb["Ground Truth Detail"], detail_df, "Ground Truth - Entity & Redaction Values (All Documents)")

    print(f"Documents: {len(doc_summaries)}")
    print(f"Ground truth rows: {len(detail_df)}")
    print(f"Categories: {len(category_df)}")
    print(f"Saved: {output_path}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--batch-root",
        default=r"output/matched_original_counterparts/matches_20260916_214743",
        help="Folder containing litellm_all_original_vs_published_redactions.json",
    )
    parser.add_argument(
        "--output",
        default=None,
        help="Output xlsx path (defaults to ground_truth_redactions.xlsx inside batch root)",
    )
    args = parser.parse_args()

    batch_root = Path(args.batch_root)
    agg_path = batch_root / "litellm_all_original_vs_published_redactions.json"
    if not agg_path.exists():
        raise SystemExit(f"Aggregate redaction JSON not found: {agg_path}")

    output_path = Path(args.output) if args.output else batch_root / "ground_truth_redactions.xlsx"
    build_workbook(agg_path, output_path)


if __name__ == "__main__":
    main()

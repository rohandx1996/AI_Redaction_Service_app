"""Compare predicted redaction JSONs with original-vs-published ground truth."""
from __future__ import annotations

import argparse
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

TITLE_FILL = PatternFill("solid", fgColor="0B2C4A")
HEADER_FILL = PatternFill("solid", fgColor="1F6FB2")
HEADER_FONT = Font(bold=True, color="FFFFFF")

FN_GROUPS = [
    "supplier name", "address", "person/signature name", "sample name",
    "product/style/model number", "report number", "PIN code", "phone/email/URL",
]
FP_GROUPS = ["chemical names", "methods", "standards", "flowcharts", "results", "units"]


def normalize(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").casefold())


def source_key(path: str) -> str:
    return Path(path or "").name.casefold()


def gt_group(category: str) -> str:
    category = category.upper()
    if category == "SUPPLIER_NAME":
        return "supplier name"
    if category == "ADDRESS":
        return "address"
    if category in {"PERSON_NAME", "APPLICANT_NAME"}:
        return "person/signature name"
    if category == "SAMPLE_NAME":
        return "sample name"
    if category == "PRODUCT_OR_PART_ID":
        return "product/style/model number"
    if category == "REPORT_OR_ACCESS_ID":
        return "report number"
    if category == "CONTACT":
        return "phone/email/URL"
    return "other sensitive entity"


def pred_group(category: str, value: str) -> str:
    category = category.upper()
    value_upper = value.upper()
    if category in {"SUPPLIER_NAME", "VENDOR_NAME", "SUBCONTRACTOR_NAME", "COMMERCIAL_PARTNER_NAME"}:
        return "supplier name"
    if category in {"ADDRESS", "LOCATION"}:
        return "address"
    if category in {"PERSON_NAME", "MANAGER_NAME", "APPROVER_NAME", "REVIEWER_NAME", "SIGNATURE", "SIGNATURE_BLOCK"}:
        return "person/signature name"
    if category in {"SAMPLE_NAME"}:
        return "sample name"
    if category in {"PART_NUMBER", "PRODUCT_CODE", "PRODUCT_NAME", "MATERIAL_ID", "MODEL_NUMBER", "PRODUCT_OR_PART_ID"}:
        return "product/style/model number"
    if category in {"REPORT_NUMBER", "DOCUMENT_ID", "CERTIFICATE_NUMBER", "PIN_CODE", "QR_CODE_TEXT", "BARCODE_TEXT", "SUPPLIER_ID", "SUPPLIER_ACCOUNT_ID", "SUPPLIER_REFERENCE", "VENDOR_ID", "TRACKING_CODE", "QR_CODE", "BARCODE", "DATAMATRIX"}:
        return "report number"
    if category in {"PHONE", "MOBILE", "FAX", "EMAIL", "URL"}:
        return "phone/email/URL"
    if category in {"CHEMICAL_NAME", "CHEMICAL", "SUBSTANCE_NAME"} or any(token in value_upper for token in {"CADMIUM", "MERCURY", "LEAD", "PHTHALATE", "HALOGEN"}):
        return "chemical names"
    if category in {"METHOD", "TEST_METHOD", "ANALYSIS_METHOD"}:
        return "methods"
    if category in {"STANDARD", "REGULATION", "DIRECTIVE"}:
        return "standards"
    if category in {"FLOWCHART", "PROCESS_STEP"}:
        return "flowcharts"
    if category in {"RESULT", "TEST_RESULT", "MEASUREMENT"}:
        return "results"
    if category in {"UNIT", "MEASUREMENT_UNIT"}:
        return "units"
    return "other extracted entity"


def load_ground_truth(path: Path) -> dict[str, list[dict]]:
    data = json.loads(path.read_text(encoding="utf-8"))
    by_source: dict[str, list[dict]] = defaultdict(list)
    for document in data.get("documents", []):
        source = source_key(document.get("original_pdf", ""))
        by_source[source]
        for item in document.get("redacted_information", []) or []:
            value = str(item.get("original_value", "")).strip()
            if value:
                by_source[source].append({
                    "document": source,
                    "page": int(item.get("page_number", 0)),
                    "category": item.get("category", ""),
                    "group": gt_group(item.get("category", "")),
                    "value": value,
                    "norm_value": normalize(value),
                    "field_label": item.get("field_label", ""),
                })
    return dict(by_source)


def load_predictions(batch_root: Path) -> dict[str, list[dict]]:
    by_source: dict[str, list[dict]] = defaultdict(list)
    for path in sorted(batch_root.glob("*/redaction_information.json")):
        data = json.loads(path.read_text(encoding="utf-8"))
        source = source_key(data.get("source_pdf", ""))
        for page_data in data.get("pages", []) or []:
            page = int(page_data.get("page_number", 0)) + 1
            for category, entities in (page_data.get("entities", {}) or {}).items():
                for entity in entities or []:
                    value = str(entity.get("value", "")).strip()
                    if value:
                        by_source[source].append({
                            "document": source,
                            "page": page,
                            "category": category,
                            "group": pred_group(category, value),
                            "value": value,
                            "norm_value": normalize(value),
                            "span_indices": entity.get("span_indices", []),
                            "confidence": entity.get("confidence"),
                        })
    return dict(by_source)


def match_entities(ground_truth: list[dict], predictions: list[dict]) -> tuple[list[dict], list[dict], list[dict]]:
    remaining = defaultdict(list)
    for index, item in enumerate(predictions):
        remaining[(item["page"], item["norm_value"])].append(index)
    matches = []
    missed = []
    matched_prediction_indices: set[int] = set()
    for truth in ground_truth:
        candidates = remaining[(truth["page"], truth["norm_value"])]
        if not candidates:
            missed.append(truth)
            continue
        prediction_index = candidates.pop(0)
        matched_prediction_indices.add(prediction_index)
        prediction = predictions[prediction_index]
        matches.append({"truth": truth, "prediction": prediction, "category_exact": truth["category"].upper() == prediction["category"].upper(), "category_group": truth["group"] == prediction["group"]})
    false_positives = [item for index, item in enumerate(predictions) if index not in matched_prediction_indices]
    return matches, missed, false_positives


def ratio(numerator: int, denominator: int) -> float:
    return numerator / denominator if denominator else 0.0


def evaluate(ground_truth: dict[str, list[dict]], predictions: dict[str, list[dict]]) -> tuple[list[dict], list[dict], list[dict], list[dict]]:
    document_rows = []
    page_rows = []
    category_rows = []
    error_rows = []
    for source in sorted(set(ground_truth) & set(predictions)):
        truth = ground_truth[source]
        pred = predictions[source]
        matches, missed, false_positives = match_entities(truth, pred)
        tp = len(matches)
        group_correct = sum(item["category_group"] for item in matches)
        exact_category = sum(item["category_exact"] for item in matches)
        prediction_key_counts = Counter((item["page"], item["norm_value"], item["category"]) for item in pred)
        duplicate_keys = [key for key, count in prediction_key_counts.items() if count > 1]
        duplicate_count = sum(prediction_key_counts[key] - 1 for key in duplicate_keys)
        document_rows.append({
            "Document": source, "Ground Truth Entities": len(truth), "Predicted Entities": len(pred), "True Positives": tp,
            "Entity Recall": ratio(tp, len(truth)), "Entity Precision": ratio(tp, len(pred)),
            "Entity F1 Score": ratio(2 * tp, len(truth) + len(pred)), "Exact-value Match Rate": ratio(tp, len(truth)),
            "Correct Span-index Rate": "N/A (ground truth has no span indices)", "Category Accuracy": ratio(group_correct, tp),
            "Exact Category Accuracy": ratio(exact_category, tp), "Duplicate Entity Rate": ratio(duplicate_count, len(pred)),
            "False-negative Rate": ratio(len(missed), len(truth)), "False-positive Rate": ratio(len(false_positives), len(pred)),
        })
        pages = sorted({item["page"] for item in truth + pred})
        for page in pages:
            page_truth = [item for item in truth if item["page"] == page]
            page_pred = [item for item in pred if item["page"] == page]
            page_matches, page_missed, page_fp = match_entities(page_truth, page_pred)
            page_rows.append({"Document": source, "Page": page, "Ground Truth Entities": len(page_truth), "Predicted Entities": len(page_pred), "True Positives": len(page_matches), "Entity Recall": ratio(len(page_matches), len(page_truth)), "Entity Precision": ratio(len(page_matches), len(page_pred)), "Entity F1 Score": ratio(2 * len(page_matches), len(page_truth) + len(page_pred)), "Exact-value Match Rate": ratio(len(page_matches), len(page_truth)), "False-negative Rate": ratio(len(page_missed), len(page_truth)), "False-positive Rate": ratio(len(page_fp), len(page_pred))})
        truth_groups = Counter(item["group"] for item in truth)
        pred_groups = Counter(item["group"] for item in pred)
        matched_groups = Counter(item["truth"]["group"] for item in matches if item["category_group"])
        for group in sorted(set(truth_groups) | set(pred_groups)):
            group_missed = sum(item["group"] == group for item in missed)
            group_fp = sum(item["group"] == group for item in false_positives)
            category_rows.append({"Document": source, "Category Group": group, "Ground Truth Count": truth_groups[group], "Predicted Count": pred_groups[group], "Matched Count": matched_groups[group], "Recall": ratio(matched_groups[group], truth_groups[group]), "Precision": ratio(matched_groups[group], pred_groups[group]), "False-negative Rate": ratio(group_missed, truth_groups[group]), "False-positive Rate": ratio(group_fp, pred_groups[group])})
        for item in missed:
            error_rows.append({"Document": source, "Page": item["page"], "Error Type": "False Negative", "Category": item["category"], "Category Group": item["group"], "Value": item["value"], "Field Label": item.get("field_label", "")})
        for item in false_positives:
            error_rows.append({"Document": source, "Page": item["page"], "Error Type": "False Positive", "Category": item["category"], "Category Group": item["group"], "Value": item["value"], "Field Label": ""})
    return document_rows, page_rows, category_rows, error_rows


def style_workbook(path: Path) -> None:
    from openpyxl import load_workbook
    workbook = load_workbook(path)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A3"
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=sheet.max_column)
        title = sheet.cell(1, 1)
        title.fill = TITLE_FILL
        title.font = Font(bold=True, color="FFFFFF", size=14)
        for cell in sheet[2]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for column_index, column_cells in enumerate(sheet.columns, start=1):
            letter = get_column_letter(column_index)
            width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 55)
            sheet.column_dimensions[letter].width = max(width, 12)
    workbook.save(path)


def build_report(gt_path: Path, batch_root: Path, output_path: Path) -> None:
    ground_truth = load_ground_truth(gt_path)
    predictions = load_predictions(batch_root)
    document_rows, page_rows, category_rows, error_rows = evaluate(ground_truth, predictions)
    missing_predictions = sorted(set(ground_truth) - set(predictions))
    compared_truth_entities = sum(len(ground_truth[k]) for k in ground_truth if k in predictions)
    full_truth_entities = sum(len(v) for v in ground_truth.values())
    predicted_entities = sum(len(predictions[k]) for k in ground_truth if k in predictions)
    true_positives = sum(r["True Positives"] for r in document_rows)
    category_group_matches = sum(r["Matched Count"] for r in category_rows)
    false_negatives = sum(r["Error Type"] == "False Negative" for r in error_rows)
    false_positives = sum(r["Error Type"] == "False Positive" for r in error_rows)
    missing_truth_entities = sum(len(ground_truth[k]) for k in missing_predictions)
    summary = pd.DataFrame([{
        "Ground Truth Documents": len(ground_truth), "Prediction Documents": len(predictions), "Compared Documents": len(document_rows),
        "Missing Prediction Documents": len(missing_predictions), "Ground Truth Entities (All)": full_truth_entities,
        "Ground Truth Entities (Compared)": compared_truth_entities, "Predicted Entities (Compared)": predicted_entities,
        "True Positives": true_positives, "Entity Recall (Compared)": ratio(true_positives, compared_truth_entities),
        "Entity Recall (All GT; Missing Docs = FN)": ratio(true_positives, full_truth_entities), "Entity Precision": ratio(true_positives, predicted_entities),
        "Entity F1 Score": ratio(2 * true_positives, compared_truth_entities + predicted_entities), "Exact-value Match Rate": ratio(true_positives, compared_truth_entities),
        "Category Accuracy": ratio(category_group_matches, true_positives), "False Negatives (Compared)": false_negatives,
        "False Negatives (Including Missing Docs)": false_negatives + missing_truth_entities, "False Positives": false_positives,
        "Span-index Rate": "N/A: ground truth JSON has no span_indices",
        "Matching Rule": "One-to-one exact normalized value on the same 1-based page; category accuracy reported separately",
    }])

    def requested_rates(groups: list[str]) -> pd.DataFrame:
        rows = []
        for group in groups:
            group_rows = [r for r in category_rows if r["Category Group"] == group]
            truth_count = sum(r["Ground Truth Count"] for r in group_rows)
            predicted_count = sum(r["Predicted Count"] for r in group_rows)
            matched_count = sum(r["Matched Count"] for r in group_rows)
            rows.append({"Category Group": group, "Ground Truth Count": truth_count, "Predicted Count": predicted_count, "Matched Count": matched_count, "Recall": ratio(matched_count, truth_count), "Precision": ratio(matched_count, predicted_count), "False-negative Rate": ratio(truth_count - matched_count, truth_count), "False-positive Rate": ratio(predicted_count - matched_count, predicted_count)})
        return pd.DataFrame(rows)

    requested_fn = requested_rates(FN_GROUPS)
    requested_fp = requested_rates(FP_GROUPS)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Summary", index=False, startrow=1)
        pd.DataFrame(document_rows).to_excel(writer, sheet_name="Document Metrics", index=False, startrow=1)
        pd.DataFrame(page_rows).to_excel(writer, sheet_name="Page Metrics", index=False, startrow=1)
        pd.DataFrame(category_rows).to_excel(writer, sheet_name="Category Metrics", index=False, startrow=1)
        requested_fn.to_excel(writer, sheet_name="Requested FN Rates", index=False, startrow=1)
        requested_fp.to_excel(writer, sheet_name="Requested FP Rates", index=False, startrow=1)
        pd.DataFrame(error_rows).to_excel(writer, sheet_name="Error Details", index=False, startrow=1)
        pd.DataFrame({"Missing Ground Truth Documents": missing_predictions}).to_excel(writer, sheet_name="Missing Predictions", index=False, startrow=1)
        for sheet_name, title in {"Summary": "Redaction Ground Truth Comparison - Summary", "Document Metrics": "Per Document Metrics", "Page Metrics": "Per Page Metrics", "Category Metrics": "Category Group Metrics", "Requested FN Rates": "False Negative Rates - Requested Sensitive Categories", "Requested FP Rates": "False Positive Rates - Requested Non-sensitive Categories", "Error Details": "False Negative and False Positive Details", "Missing Predictions": "Ground Truth Documents Without Predictions"}.items():
            writer.sheets[sheet_name].cell(1, 1, title)
    style_workbook(output_path)
    print(json.dumps({"compared_documents": len(document_rows), "ground_truth_documents": len(ground_truth), "ground_truth_entities": sum(len(v) for v in ground_truth.values()), "false_negatives": sum(r["Error Type"] == "False Negative" for r in error_rows), "false_positives": sum(r["Error Type"] == "False Positive" for r in error_rows), "output": str(output_path)}, indent=2))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--ground-truth", default=r"output/matched_original_counterparts/matches_20260916_214743/litellm_all_original_vs_published_redactions.json")
    parser.add_argument("--prediction-batch", default=r"output/timestamped_batches/batch_20260915_083545")
    parser.add_argument("--output", default=r"output/matched_original_counterparts/matches_20260916_214743/redaction_ground_truth_comparison.xlsx")
    args = parser.parse_args()
    build_report(Path(args.ground_truth), Path(args.prediction_batch), Path(args.output))


if __name__ == "__main__":
    main()
